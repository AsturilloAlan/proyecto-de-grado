"""Pruebas del reporte de transacciones atípicas (actividad 15)."""
import re
import tempfile
from datetime import date
from unittest.mock import patch

from django.contrib.auth.models import Group, User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from analisis import servicio as servicio_analisis
from analisis.models import ResultadoTransaccion
from analisis.tests import CSV_VALIDACION, paquete_de_prueba
from registros.models import (CargaArchivo, CuentaContable, EmpresaAuditada, ErrorValidacion,
                              Gestion)

from . import servicio
from .models import ReporteGenerado

MEDIA = tempfile.mkdtemp()


@override_settings(MEDIA_ROOT=MEDIA, ANALIZAR_CON_MODELO=True)
class ReporteAtipicasTests(TestCase):
    def setUp(self):
        grupo, _ = Group.objects.get_or_create(name="Auditor")
        self.auditor = User.objects.create_user(username="auditora_rep", password="Clave-Segura123")
        self.auditor.groups.add(grupo)
        self.client.login(username="auditora_rep", password="Clave-Segura123")
        self.empresa = EmpresaAuditada.objects.create(nombre="Empresa Reporte", nit="1234567")
        self.gestion = Gestion.objects.create(
            anio=2023, fecha_inicio=date(2023, 1, 1), fecha_fin=date(2023, 12, 31))
        servicio_analisis._modelo_en_memoria.update(marca=None, paquete=None)
        with patch.object(servicio_analisis, "_paquete", return_value=paquete_de_prueba()):
            self.client.post(reverse("registros:cargar"), {
                "empresa": self.empresa.id, "gestion": self.gestion.id,
                "archivo": SimpleUploadedFile("libro.csv", CSV_VALIDACION.encode("utf-8")),
            })
        self.carga = CargaArchivo.objects.latest("id")
        self.url = reverse("reportes:atipicas", args=[self.carga.id])
        self.atipicas = list(ResultadoTransaccion.objects.filter(
            registro__carga=self.carga, atipica=True).order_by("-nivel"))

    def test_resumen_con_cifras_de_materialidad(self):
        datos = servicio.reporte_atipicas([self.carga])
        self.assertEqual(datos["total_registros"], 7)
        self.assertEqual(datos["total_atipicas"], 3)
        self.assertEqual(datos["movimiento_total"], 90460)
        self.assertEqual(datos["importe_total_atipicas"], 90000)
        self.assertAlmostEqual(datos["porcentaje_importe"], 99.49, places=2)
        self.assertEqual([n["cantidad"] for n in datos["por_nivel"]], [0, 0, 3])
        self.assertEqual(datos["por_cuenta"][0]["cantidad"], 3)

    def test_listado_priorizado_de_la_mas_a_la_menos_atipica(self):
        datos = servicio.reporte_atipicas([self.carga])
        importes = [item["registro"].debe for item in datos["listado"]]
        self.assertEqual(importes, [40000, 30000, 20000])
        self.assertEqual([item["prioridad"] for item in datos["listado"]], [1, 2, 3])
        self.assertTrue(datos["listado"][0]["motivos"])

    def test_la_prioridad_no_cambia_al_filtrar(self):
        self.atipicas[2].revision = "relevante"
        self.atipicas[2].save()
        datos = servicio.reporte_atipicas([self.carga], validacion="relevante")
        self.assertEqual([item["prioridad"] for item in datos["listado"]], [3])

    def test_filtros_desconocidos_se_ignoran(self):
        datos = servicio.reporte_atipicas([self.carga], nivel="xx", motivo="<script>", validacion="zz")
        self.assertEqual(datos["filtros_texto"], "")
        self.assertEqual(len(datos["listado"]), 3)

    def test_cruce_con_avisos_de_las_reglas(self):
        fila = self.atipicas[0].registro.fila_origen
        ErrorValidacion.objects.create(carga=self.carga, fila=fila, campo="duplicado", tipo="aviso",
                                       descripcion="Posible duplicado de la fila 1")
        respuesta = self.client.get(self.url)
        self.assertEqual(respuesta.context["coinciden_con_reglas"], 1)
        self.assertContains(respuesta, "Reglas: Posible duplicado de la fila 1")

    def test_pagina_del_reporte(self):
        self.atipicas[0].revision = "relevante"
        self.atipicas[0].comentario_revision = "Sin respaldo documental"
        self.atipicas[0].revisado_por = self.auditor
        self.atipicas[0].save()
        respuesta = self.client.get(self.url)
        self.assertContains(respuesta, "Reporte de transacciones atípicas priorizadas")
        self.assertContains(respuesta, "NIT 1234567")
        self.assertContains(respuesta, "Sin respaldo documental")
        self.assertContains(respuesta, "Revisión en curso:")  # quedan 2 sin validar
        self.assertContains(respuesta, "Cómo leer este reporte")

    def test_cada_emision_queda_registrada_sin_duplicar_recargas(self):
        self.client.get(self.url)
        self.client.get(self.url)
        self.assertEqual(ReporteGenerado.objects.filter(tipo="atipicas").count(), 1)
        self.client.get(self.url, {"nivel": "medio"})
        emisiones = ReporteGenerado.objects.filter(tipo="atipicas").order_by("id")
        self.assertEqual(emisiones.count(), 2)
        self.assertEqual(emisiones.last().filtros, "Nivel medio")
        self.assertEqual(emisiones.last().total_atipicas, 3)
        self.assertEqual(emisiones.last().usuario, self.auditor)

    def test_si_cambia_la_revision_se_registra_una_emision_nueva(self):
        self.client.get(self.url)
        self.atipicas[0].revision = "relevante"
        self.atipicas[0].save()
        self.client.get(self.url)
        self.assertEqual(ReporteGenerado.objects.filter(tipo="atipicas").count(), 2)

    def test_pasar_de_pagina_no_registra_otra_emision(self):
        self.client.get(self.url)
        antes = ReporteGenerado.objects.count()
        respuesta = self.client.get(self.url, {"pagina": "2"})
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(ReporteGenerado.objects.count(), antes)
        self.assertIn("pagina_listado", respuesta.context)

    def test_volver_lleva_a_donde_estaba(self):
        url = reverse("reportes:hallazgos_gestion", args=[self.empresa.id, self.gestion.id])
        desde_carga = self.client.get(url, {"desde": f"carga-{self.carga.id}"})
        self.assertEqual(desde_carga.context["origen"]["url"], reverse("registros:detalle_carga", args=[self.carga.id]))
        self.assertContains(desde_carga, "Volver a la carga")
        desde_reportes = self.client.get(url)
        self.assertContains(desde_reportes, "Volver a Reportes")
        desde_graficos = self.client.get(url, {"desde": "graficos"})
        self.assertContains(desde_graficos, "Volver al análisis")
        otro = self.client.get(url, {"desde": "carga-999999"})
        self.assertEqual(otro.context["origen"]["tipo"], "reportes")

    def test_hallazgos_dice_que_procede(self):
        url = reverse("reportes:hallazgos_gestion", args=[self.empresa.id, self.gestion.id])
        respuesta = self.client.get(url)
        self.assertTrue(respuesta.context["que_procede"])
        self.assertContains(respuesta, respuesta.context["que_procede"][0]["texto"])

    def test_indice_muestra_la_carga_y_lo_emitido(self):
        self.client.get(self.url)
        respuesta = self.client.get(reverse("reportes:indice"))
        self.assertContains(respuesta, "Empresa Reporte")
        self.assertContains(respuesta, reverse("reportes:atipicas_gestion",
                                               args=[self.empresa.id, self.gestion.id]))
        self.assertNotContains(respuesta, "Historial de reportes emitidos")
        self.assertContains(respuesta, "Abrir hallazgos")
        self.assertContains(respuesta, reverse("reportes:hallazgos_gestion", args=[self.empresa.id, self.gestion.id]))
        self.assertNotContains(respuesta, "Comparativo PDF")  # Una sola gestión no necesita comparación.
        self.assertEqual(len(respuesta.context["emitidos"]), 1)

    def test_el_reporte_de_observaciones_tambien_se_registra(self):
        self.client.get(reverse("registros:reporte_observaciones", args=[self.carga.id]))
        self.assertTrue(ReporteGenerado.objects.filter(tipo="observaciones", carga=self.carga).exists())

    def test_sin_analisis_redirige_a_la_carga(self):
        otra = CargaArchivo.objects.create(
            archivo=SimpleUploadedFile("x.csv", b"fecha,cuenta,debe,haber\n"),
            usuario=self.auditor, empresa=self.empresa, gestion=self.gestion, estado="validado")
        respuesta = self.client.get(reverse("reportes:atipicas", args=[otra.id]))
        self.assertEqual(respuesta.status_code, 302)

    def test_requiere_sesion_y_rol(self):
        self.client.logout()
        respuesta = self.client.get(self.url)
        self.assertEqual(respuesta.status_code, 302)
        self.assertIn(reverse("login"), respuesta["Location"])

    # Organización por empresa y gestión (uso real)

    def _subir(self, gestion, contenido=CSV_VALIDACION):
        with patch.object(servicio_analisis, "_paquete", return_value=paquete_de_prueba()):
            self.client.post(reverse("registros:cargar"), {
                "empresa": self.empresa.id, "gestion": gestion.id,
                "archivo": SimpleUploadedFile("libro.csv", contenido.encode("utf-8")),
            })
        return CargaArchivo.objects.latest("id")

    def test_reporte_de_gestion_consolida_cargas_validas_y_excluye_anuladas(self):
        segunda = self._subir(self.gestion, "fecha,cuenta,glosa,debe,haber\n"
                              "05/03/2023,5101,Gasto,10,0\n06/03/2023,5101,Gasto,12,0\n"
                              "07/03/2023,5101,Gasto grande,90000,0\n")
        anulada = self._subir(self.gestion)
        CargaArchivo.objects.filter(pk=anulada.pk).update(estado="anulada")
        url = reverse("reportes:atipicas_gestion", args=[self.empresa.id, self.gestion.id])
        respuesta = self.client.get(url)
        datos = respuesta.context
        self.assertEqual({c.id for c in datos["cargas"]}, {self.carga.id, segunda.id})
        self.assertTrue(datos["varias_cargas"])
        self.assertEqual(datos["total_registros"], 10)
        self.assertEqual(datos["listado"][0]["carga_id"], segunda.id)  # 90.000 es la más atípica
        self.assertContains(respuesta, "Consolidado de 2 cargas")
        self.assertContains(respuesta, f"No: anulada")
        emision = ReporteGenerado.objects.get(tipo="atipicas", carga__isnull=True)
        self.assertEqual((emision.empresa, emision.gestion), (self.empresa, self.gestion))
        self.assertIn("consolidado", emision.alcance)

    def test_detecta_el_mismo_libro_subido_dos_veces(self):
        duplicada = self._subir(self.gestion)
        grupo = servicio.cargas_de_gestion(self.empresa, self.gestion)
        self.assertEqual(grupo["posibles_duplicadas"], [(self.carga, duplicada)])
        respuesta = self.client.get(reverse("reportes:indice"))
        self.assertContains(respuesta, "Posible carga duplicada")

    def test_indice_muestra_una_gestion_y_permite_elegir_otra(self):
        g2024 = Gestion.objects.create(
            anio=2024, fecha_inicio=date(2024, 1, 1), fecha_fin=date(2024, 12, 31))
        self._subir(g2024, CSV_VALIDACION.replace("/2023", "/2024"))
        filas = servicio.comparativo_empresa(self.empresa)
        self.assertEqual([f["gestion"].anio for f in filas], [2023, 2024])
        self.assertEqual([f["atipicas"] for f in filas], [3, 3])
        respuesta = self.client.get(reverse("reportes:indice"), {"empresa": self.empresa.id})
        gestiones = respuesta.context["expedientes"][0]["gestiones"]
        self.assertEqual([g["gestion"].anio for g in gestiones], [2024])
        self.assertEqual([g.anio for g in respuesta.context["opciones_gestion"]], [2024, 2023])
        respuesta = self.client.get(reverse("reportes:indice"), {
            "empresa": self.empresa.id, "gestion": self.gestion.id})
        self.assertEqual([g["gestion"].anio for g in respuesta.context["expedientes"][0]["gestiones"]], [2023])
        respuesta = self.client.get(reverse("reportes:indice"), {
            "empresa": self.empresa.id, "gestion": "999999"})
        self.assertEqual(respuesta.context["gestion_id"], g2024.id)

    def test_indice_no_mezcla_empresas_al_seleccionar(self):
        otra = EmpresaAuditada.objects.create(nombre="Otra empresa", nit="7654321")
        CargaArchivo.objects.create(empresa=otra, gestion=self.gestion,
                                   usuario=self.auditor, estado="validado", archivo="demo.csv")
        respuesta = self.client.get(reverse("reportes:indice"), {"empresa": otra.id})
        self.assertEqual(len(respuesta.context["expedientes"]), 1)
        self.assertEqual(respuesta.context["expedientes"][0]["empresa"], otra)
        respuesta = self.client.get(reverse("reportes:indice"), {"empresa": "999999"})
        self.assertEqual(respuesta.context["expedientes"], [])

    def test_reporte_de_observaciones_de_la_gestion(self):
        fila = self.atipicas[0].registro.fila_origen
        ErrorValidacion.objects.create(
            carga=self.carga, fila=fila, campo="duplicado", tipo="aviso", revisado=True,
            estado_revision="observado", descripcion="Posible duplicado de la fila 1",
            comentario_revision="Pago registrado dos veces", revisado_por=self.auditor)
        respuesta = self.client.get(
            reverse("reportes:observaciones_gestion", args=[self.empresa.id, self.gestion.id]))
        self.assertContains(respuesta, "Pago registrado dos veces")
        self.assertContains(respuesta, "Avisos por tipo de regla")
        self.assertContains(respuesta, "Posible duplicado")
        self.assertEqual(respuesta.context["importe_observado"], 40000)
        self.assertTrue(ReporteGenerado.objects.filter(tipo="observaciones", carga__isnull=True).exists())

    # Cuentas tributarias, hallazgos y variación entre gestiones

    def test_reconoce_cuentas_tributarias_por_su_nombre(self):
        for nombre in ("DEBITO FISCAL - IVA", "I.T. POR PAGAR", "Retenciones RC-IVA",
                       "Impuesto a las Utilidades (IUE)", "CREDITO FISCAL"):
            self.assertTrue(servicio.es_cuenta_tributaria(nombre), nombre)
        for nombre in ("CAJA MONEDA NACIONAL", "UTILIDADES ACUMULADAS", "VITALICIA SEGUROS",
                       "DEPRECIACIÓN ACUMULADA"):
            self.assertFalse(servicio.es_cuenta_tributaria(nombre), nombre)

    def test_seccion_y_filtro_de_cuentas_tributarias(self):
        CuentaContable.objects.filter(pk=self.atipicas[0].registro.cuenta_id).update(
            nombre="DEBITO FISCAL - IVA")
        datos = servicio.reporte_atipicas([self.carga], tributarias="1")
        self.assertEqual(datos["tributarias"]["cantidad"], 3)
        self.assertEqual(datos["tributarias"]["por_cuenta"][0]["nombre"], "DEBITO FISCAL - IVA")
        self.assertTrue(all(item["tributaria"] for item in datos["listado"]))
        self.assertEqual(datos["filtros_texto"], "Solo cuentas tributarias")
        respuesta = self.client.get(self.url, {"tributarias": "1"})
        self.assertContains(respuesta, "En cuentas de impuestos")
        # La mención al Informe Tributario Complementario quedó solo en el PDF y el Excel.
        self.assertContains(respuesta, "Solo cuentas de impuestos")

    def test_hallazgos_une_modelo_y_reglas_sin_repetir(self):
        principal = self.atipicas[0]  # 40.000
        principal.revision = "relevante"
        principal.comentario_revision = "Sin respaldo"
        principal.revisado_por = self.auditor
        principal.save()
        # La misma fila también fue observada por una regla.
        ErrorValidacion.objects.create(
            carga=self.carga, fila=principal.registro.fila_origen, campo="duplicado", tipo="aviso",
            revisado=True, estado_revision="observado", descripcion="Posible duplicado de la fila 2",
            comentario_revision="Duplicado confirmado", revisado_por=self.auditor)
        # Otra fila, que el modelo no señaló, observada por una regla.
        normal = ResultadoTransaccion.objects.filter(registro__carga=self.carga, atipica=False).first()
        ErrorValidacion.objects.create(
            carga=self.carga, fila=normal.registro.fila_origen, campo="comprobante", tipo="aviso",
            revisado=True, estado_revision="observado", descripcion="Sin número de comprobante",
            comentario_revision="Falta comprobante", revisado_por=self.auditor)
        datos = servicio.reporte_hallazgos([self.carga])
        self.assertEqual(len(datos["hallazgos"]), 2)
        self.assertEqual([h["origen"] for h in datos["hallazgos"]], ["Modelo y reglas", "Reglas"])
        self.assertEqual(datos["hallazgos"][0]["importe"], 40000)  # ordenados por importe
        self.assertEqual(datos["por_origen"], {"modelo": 0, "reglas": 1, "ambos": 1})
        self.assertTrue(datos["es_borrador"])  # quedan atípicas sin revisar
        respuesta = self.client.get(reverse("reportes:hallazgos_gestion",
                                            args=[self.empresa.id, self.gestion.id]))
        self.assertContains(respuesta, "Hallazgos de auditoría")
        self.assertContains(respuesta, "Duplicado confirmado")
        self.assertContains(respuesta, "Sin respaldo")
        self.assertTrue(ReporteGenerado.objects.filter(tipo="hallazgos").exists())

    def test_variacion_frente_a_la_gestion_anterior(self):
        g2024 = Gestion.objects.create(
            anio=2024, fecha_inicio=date(2024, 1, 1), fecha_fin=date(2024, 12, 31))
        # 2024: un libro distinto con una sola transacción grande.
        self._subir(g2024, "fecha,cuenta,glosa,debe,haber\n05/03/2024,5101,Gasto,10,0\n"
                           "06/03/2024,5101,Gasto,12,0\n07/03/2024,5101,Gasto,90000,0\n")
        filas = servicio.comparativo_empresa(self.empresa)
        self.assertIsNone(filas[0]["variacion_atipicas"])  # 2023 no tiene anterior
        self.assertEqual(filas[1]["anterior"], self.gestion)
        esperado = (filas[1]["atipicas"] - filas[0]["atipicas"]) * 100 / filas[0]["atipicas"]
        self.assertAlmostEqual(filas[1]["variacion_atipicas"], esperado)
        respuesta = self.client.get(reverse("reportes:indice"))
        self.assertNotContains(respuesta, "Complemento: comparación entre gestiones")
        self.assertNotContains(respuesta, reverse("reportes:comparativo", args=[self.empresa.id]))


@override_settings(MEDIA_ROOT=MEDIA, ANALIZAR_CON_MODELO=True)
class ExportacionTests(TestCase):
    """Fila 17: los reportes se descargan en PDF y Excel con el mismo contenido."""

    def setUp(self):
        grupo, _ = Group.objects.get_or_create(name="Auditor")
        self.auditor = User.objects.create_user(username="aud_export", password="Clave-Segura123")
        self.auditor.groups.add(grupo)
        self.client.login(username="aud_export", password="Clave-Segura123")
        self.empresa = EmpresaAuditada.objects.create(nombre="Empresa Export S.A.", nit="999")
        self.gestion = Gestion.objects.create(
            anio=2023, fecha_inicio=date(2023, 1, 1), fecha_fin=date(2023, 12, 31))
        # Una glosa que empieza como fórmula de Excel no debe ejecutarse al abrir el archivo.
        contenido = CSV_VALIDACION.replace("Pago extraordinario 5", "=SUMA(1+1)")
        servicio_analisis._modelo_en_memoria.update(marca=None, paquete=None)
        with patch.object(servicio_analisis, "_paquete", return_value=paquete_de_prueba()):
            self.client.post(reverse("registros:cargar"), {
                "empresa": self.empresa.id, "gestion": self.gestion.id,
                "archivo": SimpleUploadedFile("libro.csv", contenido.encode("utf-8")),
            })
        self.carga = CargaArchivo.objects.latest("id")
        self.url = reverse("reportes:atipicas_gestion", args=[self.empresa.id, self.gestion.id])

    def _excel(self, respuesta):
        from io import BytesIO

        from openpyxl import load_workbook

        return load_workbook(BytesIO(respuesta.content))

    def test_atipicas_en_excel(self):
        respuesta = self.client.get(self.url, {"formato": "excel"})
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("attachment;", respuesta["Content-Disposition"])
        self.assertIn("atipicas_empresa-export-sa_2023_", respuesta["Content-Disposition"])
        libro = self._excel(respuesta)
        self.assertEqual(libro.sheetnames[0], "Resumen")
        self.assertEqual(libro.sheetnames, ["Resumen", "Alcance", "Indicadores", "Validación",
                                            "Utilidad del modelo", "Listado", "Tributarias", "Por nivel", "Por motivo", "Por cuenta", "Por mes"])
        listado = libro["Listado"]
        self.assertEqual(listado.max_row, 3 + 3)  # título, vacío, cabecera y 3 filas
        valores = [c.value for fila in listado.iter_rows() for c in fila]
        formula = next(v for v in valores if isinstance(v, str) and v.startswith("=SUMA"))
        celda = next(c for fila in listado.iter_rows() for c in fila if c.value == formula)
        self.assertEqual(celda.data_type, "s")  # guardado como texto, no como fórmula

    def test_atipicas_en_pdf_respeta_los_filtros(self):
        respuesta = self.client.get(self.url, {"formato": "pdf", "nivel": "medio"})
        self.assertEqual(respuesta["Content-Type"], "application/pdf")
        self.assertTrue(respuesta.content.startswith(b"%PDF"))
        emision = ReporteGenerado.objects.get(formato="pdf")
        self.assertEqual(emision.filtros, "Nivel medio")
        self.assertEqual(emision.tipo, "atipicas")

    def test_observaciones_y_hallazgos_en_ambos_formatos(self):
        urls = [
            reverse("reportes:observaciones_gestion", args=[self.empresa.id, self.gestion.id]),
            reverse("registros:reporte_observaciones", args=[self.carga.id]),
            reverse("reportes:hallazgos_gestion", args=[self.empresa.id, self.gestion.id]),
            reverse("reportes:hallazgos", args=[self.carga.id]),
        ]
        for url in urls:
            pdf = self.client.get(url, {"formato": "pdf"})
            self.assertTrue(pdf.content.startswith(b"%PDF"), url)
            excel = self._excel(self.client.get(url, {"formato": "excel"}))
            self.assertEqual(excel.sheetnames[0], "Resumen", url)
        self.assertEqual(ReporteGenerado.objects.filter(formato="excel").count(), 4)

    def test_comparativo_de_la_empresa(self):
        url = reverse("reportes:comparativo", args=[self.empresa.id])
        libro = self._excel(self.client.get(url, {"formato": "excel"}))
        gestiones = libro["Gestiones"]
        self.assertEqual(gestiones.cell(row=4, column=1).value, 2023)
        self.assertTrue(self.client.get(url, {"formato": "pdf"}).content.startswith(b"%PDF"))
        self.assertEqual(self.client.get(url).status_code, 302)  # sin formato vuelve al índice

    def _texto_pdf(self, contenido):
        from io import BytesIO

        import pdfplumber
        with pdfplumber.open(BytesIO(contenido)) as pdf:
            return "\n".join(pagina.extract_text() or "" for pagina in pdf.pages)

    def test_firmas_de_hallazgos_segun_quien_emite(self):
        url = reverse("reportes:hallazgos_gestion", args=[self.empresa.id, self.gestion.id])
        auditor = self._texto_pdf(self.client.get(url, {"formato": "pdf"}).content)
        self.assertIn("Elaborado por", auditor)
        self.assertIn("Revisado y aprobado por", auditor)
        # El socio (Administrador) firma una sola vez.
        socio = User.objects.create_user(username="socio", password="Clave-Segura123")
        socio.groups.add(Group.objects.get_or_create(name="Administrador")[0])
        self.client.force_login(socio)
        texto = self._texto_pdf(self.client.get(url, {"formato": "pdf"}).content)
        self.assertIn("Elaborado y aprobado por", texto)
        self.assertNotIn("Revisado y aprobado por", texto)

    def test_listado_de_atipicas_no_lleva_firmas(self):
        texto = self._texto_pdf(self.client.get(self.url, {"formato": "pdf"}).content)
        self.assertNotIn("Elaborado", texto)
        self.assertIn("REPORTE DE APOYO", texto)

    def test_la_utilidad_del_modelo_va_solo_al_excel(self):
        from io import BytesIO

        from openpyxl import load_workbook

        texto = self._texto_pdf(self.client.get(self.url, {"formato": "pdf"}).content)
        self.assertNotIn("Útiles", texto)
        libro = load_workbook(BytesIO(self.client.get(self.url, {"formato": "excel"}).content))
        self.assertIn("Utilidad del modelo", libro.sheetnames)

    def test_la_pagina_muestra_los_botones_de_descarga(self):
        respuesta = self.client.get(self.url, {"nivel": "medio"})
        self.assertContains(respuesta, "nivel=medio&amp;formato=pdf")
        self.assertContains(respuesta, "Descargar Excel")

    def test_formato_desconocido_muestra_la_pagina(self):
        respuesta = self.client.get(self.url, {"formato": "exe"})
        self.assertEqual(respuesta["Content-Type"], "text/html; charset=utf-8")


@override_settings(MEDIA_ROOT=MEDIA, ANALIZAR_CON_MODELO=True)
class PanoramaTests(TestCase):
    """Fila 16: gráficos de la gestión con filtros por fecha, cuenta y nivel."""

    def setUp(self):
        grupo, _ = Group.objects.get_or_create(name="Auditor")
        usuario = User.objects.create_user(username="aud_graf", password="Clave-Segura123")
        usuario.groups.add(grupo)
        self.client.login(username="aud_graf", password="Clave-Segura123")
        self.empresa = EmpresaAuditada.objects.create(nombre="Empresa Gráficos")
        self.gestion = Gestion.objects.create(
            anio=2023, fecha_inicio=date(2023, 1, 1), fecha_fin=date(2023, 12, 31))
        contenido = ("fecha,cuenta,glosa,debe,haber\n"
                     "02/01/2023,1101,Pago,100,0\n03/01/2023,1101,Pago,110,0\n04/01/2023,1101,Pago,120,0\n"
                     "05/03/2023,1101,Pago grande,30000,0\n06/03/2023,5101,Gasto,10,0\n07/03/2023,5101,Gasto,12,0\n"
                     "08/03/2023,5101,Gasto,15,0\n09/04/2023,5101,Gasto grande,40000,0\n")
        servicio_analisis._modelo_en_memoria.update(marca=None, paquete=None)
        with patch.object(servicio_analisis, "_paquete", return_value=paquete_de_prueba()):
            self.client.post(reverse("registros:cargar"), {
                "empresa": self.empresa.id, "gestion": self.gestion.id,
                "archivo": SimpleUploadedFile("libro.csv", contenido.encode("utf-8")),
            })
        self.datos_url = reverse("reportes:panorama_datos", args=[self.empresa.id, self.gestion.id])

    def test_la_pagina_carga_al_instante_con_siluetas(self):
        respuesta = self.client.get(reverse("reportes:panorama", args=[self.empresa.id, self.gestion.id]))
        self.assertContains(respuesta, "esqueleto-grafico")
        self.assertContains(respuesta, "¿En qué meses aparecen?")
        self.assertContains(respuesta, "Ver los datos")
        self.assertContains(respuesta, self.datos_url)

    def test_datos_de_los_graficos(self):
        datos = self.client.get(self.datos_url).json()
        self.assertEqual(datos["resumen"]["atipicas"], 2)
        self.assertEqual(datos["resumen"]["importe"], 70000.0)
        self.assertEqual(datos["por_mes"]["etiquetas"], ["Mar 2023", "Abr 2023"])
        self.assertEqual(datos["por_mes"]["cantidades"], [1, 1])
        self.assertEqual(sum(datos["por_nivel"]["cantidades"]), 2)
        self.assertEqual(datos["revision"]["cantidades"], [0, 0, 0, 2])
        self.assertEqual(len(datos["opciones_cuenta"]), 2)
        self.assertEqual(datos["mas_atipicas"][0]["importe"], 40000.0)

    def test_filtros_por_fecha_cuenta_y_nivel(self):
        datos = self.client.get(self.datos_url, {"desde": "2023-04-01"}).json()
        self.assertEqual(datos["resumen"]["atipicas"], 1)
        cuenta = next(o["id"] for o in datos["opciones_cuenta"] if o["texto"].startswith("1101"))
        datos = self.client.get(self.datos_url, {"cuenta": cuenta}).json()
        self.assertEqual(datos["mas_atipicas"][0]["cuenta"][:4], "1101")
        self.assertEqual(len(datos["opciones_cuenta"]), 2)  # el selector conserva todas las cuentas
        datos = self.client.get(self.datos_url, {"nivel": "muy_alto"}).json()
        self.assertEqual(datos["resumen"]["atipicas"], 0)

    def test_filtros_invalidos_se_ignoran(self):
        datos = self.client.get(self.datos_url, {"desde": "no-fecha", "cuenta": "1 OR 1=1",
                                                 "nivel": "<b>"}).json()
        self.assertEqual(datos["resumen"]["atipicas"], 2)
        self.assertEqual(datos["filtros"], {"desde": None, "hasta": None, "cuenta": "", "nivel": ""})

    def test_requiere_rol(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.datos_url).status_code, 302)


class RangoFechasFiltroTests(TestCase):
    def setUp(self):
        from registros.models import Gestion
        self.gestion = Gestion(anio=2023, fecha_inicio=date(2023, 1, 1), fecha_fin=date(2023, 12, 31))

    def test_fechas_fuera_de_la_gestion_se_ignoran(self):
        self.assertEqual(servicio.rango_dentro_de_gestion(self.gestion, "1900-01-01", "2099-05-01"), ("", ""))

    def test_fecha_mal_escrita_se_ignora(self):
        self.assertEqual(servicio.rango_dentro_de_gestion(self.gestion, "2023-13-40", ""), ("", ""))

    def test_desde_mayor_que_hasta_se_invierte(self):
        self.assertEqual(servicio.rango_dentro_de_gestion(self.gestion, "2023-09-01", "2023-03-01"),
                         ("2023-03-01", "2023-09-01"))


class PeriodosGestionTests(TestCase):
    def test_trimestres_y_cierre_de_gestion_calendario(self):
        from registros.models import Gestion
        from reportes.views import _periodos_de_gestion
        g = Gestion(anio=2023, fecha_inicio=date(2023, 1, 1), fecha_fin=date(2023, 12, 31))
        p = {x["clave"]: (x["desde"], x["hasta"]) for x in _periodos_de_gestion(g)}
        self.assertEqual(p["t1"], ("2023-01-01", "2023-03-31"))
        self.assertEqual(p["t4"], ("2023-10-01", "2023-12-31"))
        self.assertEqual(p["cierre"], ("2023-12-01", "2023-12-31"))

    def test_gestion_fiscal_abril_marzo(self):
        from registros.models import Gestion
        from reportes.views import _periodos_de_gestion
        g = Gestion(anio=2024, fecha_inicio=date(2023, 4, 1), fecha_fin=date(2024, 3, 31))
        p = {x["clave"]: (x["desde"], x["hasta"]) for x in _periodos_de_gestion(g)}
        self.assertEqual(p["t1"], ("2023-04-01", "2023-06-30"))
        self.assertEqual(p["t4"], ("2024-01-01", "2024-03-31"))
        self.assertEqual(p["cierre"], ("2024-03-01", "2024-03-31"))



class SelectorGestionesTests(TestCase):
    def test_con_muchas_gestiones_muestra_cuatro_y_agrupa_las_anteriores(self):
        from datetime import date as d
        from registros.models import CargaArchivo, EmpresaAuditada, Gestion
        usuario = User.objects.create_user(username="aud_sel", password="x")
        usuario.groups.add(Group.objects.get_or_create(name="Auditor")[0])
        self.client.force_login(usuario)
        empresa = EmpresaAuditada.objects.create(nombre="Muchas Gestiones SRL")
        gestiones = []
        for anio in range(2016, 2025):
            g = Gestion.objects.create(anio=anio, fecha_inicio=d(anio, 1, 1), fecha_fin=d(anio, 12, 31))
            CargaArchivo.objects.create(empresa=empresa, gestion=g, usuario=usuario, archivo="x.csv", estado="validado")
            gestiones.append(g)
        respuesta = self.client.get(reverse("reportes:indice"), {"empresa": empresa.id})
        self.assertContains(respuesta, "Anteriores")
        contenido = respuesta.content.decode()
        visibles = contenido[contenido.index('class="selector-gestiones"'):contenido.index("Anteriores")]
        self.assertIn(">2024<", visibles)
        self.assertNotIn(">2016<", visibles)
        antigua = self.client.get(reverse("reportes:indice"), {"empresa": empresa.id, "gestion": gestiones[0].id})
        self.assertContains(antigua, 'class="activa"')
