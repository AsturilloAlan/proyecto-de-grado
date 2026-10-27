"""Exportación de reportes a Excel y PDF (actividad 17).

Cada reporte se describe primero como un `Documento` neutro (título, datos de
encabezado y secciones con tablas). El mismo documento se convierte a Excel o a PDF, así
la pantalla, el Excel y el PDF dicen exactamente lo mismo.
"""
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from django.utils.text import slugify

RAZON_SOCIAL = "ST&S AUDITORES Y CONSULTORES S.R.L."
PIE = ("Documento generado automáticamente por el sistema de apoyo a la auditoría externa de "
       "ST&S Auditores y Consultores S.R.L. Uso interno, sujeto a revisión profesional.")


@dataclass
class Seccion:
    titulo: str
    columnas: list
    filas: list
    # Proporción de ancho de cada columna (suma libre); None = repartir igual.
    anchos: list = None
    # Índices de columnas numéricas (alineadas a la derecha, con formato de importe).
    numericas: tuple = ()
    nota: str = ""
    # Nombre corto de la hoja de Excel (máximo 31 caracteres).
    hoja: str = ""
    # El detalle estadístico se incluye solo en Excel.
    solo_excel: bool = False


@dataclass
class Documento:
    tipo: str
    titulo: str
    subtitulo: str
    datos: list  # [(etiqueta, valor)]
    secciones: list = field(default_factory=list)
    avisos: list = field(default_factory=list)  # borrador, carga anulada, duplicadas
    emision: int = None
    nombre_archivo: str = "reporte"
    emitido_por: str = ""
    emisor_es_socio: bool = False
    # Reporte de observaciones: lo recibe el representante legal de la empresa auditada.
    firma_cliente: str = ""
    # Contador de la empresa (firma la justificación de las observaciones).
    contador: str = ""


# Formato de valores

def importe(valor):
    """1234567.8 → '1.234.567,80' (formato boliviano)."""
    if valor in (None, ""):
        return ""
    texto = f"{Decimal(valor):,.2f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def porcentaje(valor, decimales=1):
    if valor is None:
        return "-"
    return f"{valor:.{decimales}f} %".replace(".", ",")


def fecha(valor, con_hora=False):
    if not valor:
        return ""
    if hasattr(valor, "tzinfo") and con_hora:
        valor = timezone.localtime(valor)
        return valor.strftime("%d/%m/%Y %H:%M")
    return valor.strftime("%d/%m/%Y")


def _importe_lado(registro):
    if registro is None:
        return ""
    return f"{importe(registro.debe)} D" if registro.debe else f"{importe(registro.haber)} H"


def _nombre_archivo(tipo, empresa, gestion, carga):
    partes = [tipo, slugify(empresa.nombre)[:40], str(gestion.anio)]
    if carga is not None:
        partes.append(f"carga{carga.id}")
    partes.append(timezone.localtime().strftime("%Y%m%d_%H%M"))
    return "_".join(partes)


def _datos_comunes(contexto):
    empresa, gestion, carga = contexto["empresa"], contexto["gestion"], contexto.get("carga")
    datos = [
        ("Empresa auditada", f"{empresa}" + (f" (NIT {empresa.nit})" if empresa.nit else "")),
        ("Gestión", str(gestion)),
        ("Alcance", f"Carga N.º {carga.id}" if carga else
         f"Consolidado de {len(contexto.get('cargas', []))} carga(s) válida(s) de la gestión"),
        ("Emitido", f"{fecha(contexto['fecha_generacion'], True)} por {contexto['generado_por'].username}"),
        ("Emisión N.º", str(contexto["emision"].id)),
    ]
    return datos


def _avisos_comunes(contexto):
    avisos = []
    carga = contexto.get("carga")
    if carga is not None and carga.estado == "anulada":
        avisos.append("CARGA ANULADA: este reporte no tiene validez como papel de trabajo.")
    grupo = contexto.get("grupo")
    if grupo:
        for a, b in grupo["posibles_duplicadas"]:
            avisos.append(f"Posible carga duplicada: N.º {a.id} y N.º {b.id} parecen el mismo libro.")
    return avisos


def _seccion_alcance(contexto, sin_analisis=()):
    grupo = contexto.get("grupo")
    if not grupo:
        return None
    filas = []
    for c in grupo["incluidas"]:
        filas.append([f"N.º {c.id}", c.get_formato_detectado_display() or "Sin dato",
                      f"{fecha(c.fecha_carga)} · {c.usuario.username}", c.registros_validos,
                      "No: sin análisis del modelo" if c in sin_analisis else "Sí"])
    for c in grupo["anuladas"]:
        filas.append([f"N.º {c.id}", c.get_formato_detectado_display() or "Sin dato",
                      f"{fecha(c.fecha_carga)} · {c.usuario.username}", c.registros_validos,
                      f"No: anulada el {fecha(c.fecha_anulacion)}"])
    for c in grupo["en_proceso"]:
        filas.append([f"N.º {c.id}", "", "", "", "No: en proceso"])
    return Seccion(f"Alcance: cargas de la gestión {contexto['gestion'].anio}",
                   ["Carga", "Tipo de libro", "Subida", "Filas guardadas", "En este reporte"],
                   filas, anchos=[1, 2, 2, 1.2, 2.5], hoja="Alcance")


# Documentos

def documento_atipicas(d):
    v = d["validacion"]
    datos = _datos_comunes(d) + [
        ("Modelo aplicado", "; ".join(sorted({f"{a.algoritmo} ({a.configuracion})" for a in d["analisis"]}))),
        ("Filtros del listado", d["filtros_texto"] or "Sin filtros"),
    ]
    avisos = _avisos_comunes(d)
    if d["es_borrador"]:
        avisos.append(f"VERSIÓN PRELIMINAR: revisión en curso, {v['revisadas']} de {d['total_atipicas']} "
                      "atípicas revisadas.")

    secciones = []
    alcance = _seccion_alcance(d, d["sin_analisis"])
    if alcance:
        secciones.append(alcance)
    secciones.append(Seccion("1. Resumen", ["Indicador", "Valor"], [
        ["Registros analizados", f"{d['total_registros']:,}".replace(",", ".")],
        ["Transacciones atípicas", f"{d['total_atipicas']} ({porcentaje(d['porcentaje_atipicas'], 2)} de los registros)"],
        ["Importe involucrado", importe(d["importe_total_atipicas"])],
        ["   Debe / Haber", f"{importe(d['importe_atipicas']['debe'])} / {importe(d['importe_atipicas']['haber'])}"],
        ["% del movimiento total", f"{porcentaje(d['porcentaje_importe'], 2)} de {importe(d['movimiento_total'])}"],
        ["También señaladas por las reglas", str(d["coinciden_con_reglas"])],
    ], anchos=[2, 3], hoja="Indicadores"))
    secciones.append(Seccion("2. Validación del auditor",
        ["Revisadas", "Pendientes", "Observadas", "En seguimiento", "Válidas"],
        [[f"{v['revisadas']} de {v['atipicas']}", v["pendientes"], v["observadas"], v["seguimiento"],
          v["no_relevantes"]]], hoja="Validación"))
    # Utilidad del modelo (validación de la investigación): solo en el Excel, no en el documento.
    secciones.append(Seccion("Utilidad de las señales del modelo",
        ["Útiles sobre lo revisado", f"Útiles entre las {v['primeras']} primeras"],
        [[porcentaje(v["precision"]), porcentaje(v["precision_primeras"])]],
        nota="Útil = observada o en seguimiento.", hoja="Utilidad del modelo", solo_excel=True))
    columnas = ["N.º", "Nivel", "Fecha", "Fila", "Comprobante", "Cuenta", "Glosa", "Importe",
                "Por qué la señala el modelo", "Reglas", "Validación", "Comentario", "Revisado por"]
    anchos = [0.55, 1.05, 0.95, 0.65, 1.3, 1.9, 2.3, 1.35, 3.0, 1.5, 1.2, 1.9, 1.2]
    if d["varias_cargas"]:
        columnas.insert(1, "Carga")
        anchos.insert(1, 0.6)
    filas = []
    for item in d["listado"]:
        r, res = item["registro"], item["resultado"]
        fila = [item["prioridad"], f"{res.nivel_texto} ({res.nivel:.2f})".replace(".", ","),
                fecha(r.fecha), r.fila_origen, r.numero_comprobante,
                f"{r.cuenta.codigo} {r.cuenta.nombre}" + (" [tributaria]" if item["tributaria"] else ""),
                r.glosa, _importe_lado(r),
                "\n".join(("También: " if i else "") + m["texto"] for i, m in enumerate(item["motivos"])),
                "; ".join(a["descripcion"] for a in item["avisos"]),
                res.get_revision_display(), res.comentario_revision,
                f"{res.revisado_por.username if res.revisado_por else ''} {fecha(res.fecha_revision)}".strip()]
        if d["varias_cargas"]:
            fila.insert(1, item["carga_id"])
        filas.append(fila)
    secciones.append(Seccion(f"3. Listado priorizado ({len(filas)} transacciones)", columnas, filas,
                             anchos=anchos, numericas=(columnas.index("Importe"),), hoja="Listado"))
    t = d["tributarias"]
    secciones.append(Seccion("4. Atípicas en cuentas tributarias",
        ["Cuenta", "Nombre", "Atípicas", "Observadas", "Importe"],
        [[c["codigo"], c["nombre"], c["cantidad"], c["relevantes"], importe(c["importe"])] for c in t["por_cuenta"]],
        anchos=[1.3, 3.5, 1, 1, 1.5], numericas=(2, 3, 4),
        nota=(f"{t['cantidad']} atípicas por {importe(t['importe'])}: {t['relevantes']} observadas y "
              f"{t['pendientes']} pendientes. Cuentas de IVA, IT, RC-IVA, IUE, ICE, retenciones o "
              "impuestos; apoya la revisión para el Informe Tributario Complementario."),
        hoja="Tributarias"))
    secciones.append(Seccion("Distribución por nivel de anomalía",
        ["Nivel", "Cantidad", "% de las atípicas", "Importe"],
        [[n["etiqueta"], n["cantidad"], porcentaje(n["porcentaje"]), importe(n["importe"])] for n in d["por_nivel"]],
        numericas=(1, 3), hoja="Por nivel", solo_excel=True))
    secciones.append(Seccion("Distribución por motivo principal",
        ["Motivo principal", "Cantidad", "% de las atípicas", "Importe"],
        [[m["etiqueta"], m["cantidad"], porcentaje(m["porcentaje"]), importe(m["importe"])] for m in d["por_motivo"]],
        anchos=[3, 1, 1.3, 1.5], numericas=(1, 3), hoja="Por motivo", solo_excel=True))
    secciones.append(Seccion("Cuentas con más transacciones atípicas",
        ["Cuenta", "Nombre", "Grupo", "Atípicas", "Observadas", "Importe"],
        [[c["codigo"], c["nombre"], (c["tipo"] or "").capitalize(), c["cantidad"], c["relevantes"],
          importe(c["importe"])] for c in d["por_cuenta"]],
        anchos=[1.3, 3, 1.1, 0.9, 0.9, 1.5], numericas=(3, 4, 5),
        nota="Hasta 15 cuentas, de la que más atípicas tiene a la que menos.", hoja="Por cuenta",
        solo_excel=True))
    secciones.append(Seccion("Distribución por mes", ["Mes", "Atípicas", "Importe"],
        [[f["mes"].strftime("%m/%Y") if f["mes"] else "", f["cantidad"], importe(f["importe"])] for f in d["por_mes"]],
        numericas=(1, 2), hoja="Por mes", solo_excel=True))

    return Documento(
        tipo="atipicas", titulo="Reporte de transacciones atípicas priorizadas",
        subtitulo=("Transacciones que el modelo de Machine Learning señaló como atípicas, de la más a "
                   "la menos atípica, con su motivo y la validación del auditor. El modelo señala casos "
                   "para revisar; la conclusión corresponde al criterio profesional del auditor."),
        datos=datos, secciones=secciones, avisos=avisos, emision=d["emision"].id,
        nombre_archivo=_nombre_archivo("atipicas", d["empresa"], d["gestion"], d.get("carga")))


def documento_observaciones(d):
    avisos = _avisos_comunes(d)
    if d["avisos_pendientes"]:
        avisos.append(f"VERSIÓN PRELIMINAR: faltan revisar {d['avisos_pendientes']} de {d['avisos_total']} avisos.")
    datos = _datos_comunes(d) + [
        ("Observaciones", str(len(d["filas_reporte"]))),
        ("Importe observado", importe(d["importe_observado"])),
        ("Cobertura", f"{d['avisos_total']} avisos: {len(d['filas_reporte'])} observados, "
                      f"{d['avisos_validos']} válidos, {d['avisos_pendientes']} pendientes"),
    ]
    secciones = []
    alcance = _seccion_alcance(d)
    if alcance:
        secciones.append(alcance)
    secciones.append(Seccion("Avisos por tipo de regla",
        ["Regla", "Avisos", "Observados", "Válidos", "Pendientes"],
        [[t["etiqueta"], t["total"], t["observados"], t["validos"], t["pendientes"]] for t in d["por_tipo"]],
        # Conteo interno de la revisión: el documento que firma la empresa trae solo las observaciones.
        anchos=[3.5, 1, 1, 1, 1], numericas=(1, 2, 3, 4), hoja="Por regla", solo_excel=True))
    # "Justificación del contador" va vacía: la escribe el contador de la empresa al recibir el reporte.
    columnas = ["Fila", "Fecha", "Comprobante", "Cuenta", "Importe", "Hallazgo", "Glosa",
                "Fundamento del auditor", "Revisado por", "Justificación del contador"]
    anchos = [0.6, 0.95, 1.2, 1.9, 1.25, 2.1, 1.9, 2.2, 1.1, 2.8]
    if d["varias_cargas"]:
        columnas.insert(0, "Carga")
        anchos.insert(0, 0.6)
    filas = []
    for item in d["filas_reporte"]:
        o, r = item["observacion"], item["registro"]
        fila = [o.fila, fecha(r.fecha) if r else "", r.numero_comprobante if r else "",
                f"{r.cuenta.codigo} {r.cuenta.nombre}" if r else "",
                _importe_lado(r), o.descripcion, r.glosa if r else "", o.comentario_revision,
                f"{o.revisado_por.username if o.revisado_por else ''} {fecha(o.fecha_revision)}".strip(), "\n\n\n"]
        if d["varias_cargas"]:
            fila.insert(0, item["carga_id"])
        filas.append(fila)
    secciones.append(Seccion("Detalle de las observaciones", columnas, filas, anchos=anchos,
                             numericas=(columnas.index("Importe"),), hoja="Observaciones"))
    return Documento(
        tipo="observaciones", titulo="Reporte de transacciones observadas",
        subtitulo="Hallazgos confirmados por el auditor sobre los avisos de las reglas de validación.",
        datos=datos, secciones=secciones, avisos=avisos, emision=d["emision"].id,
        nombre_archivo=_nombre_archivo("observaciones", d["empresa"], d["gestion"], d.get("carga")),
        contador=d["empresa"].contador_nombre,
        firma_cliente=(f"Recibido por: {d['empresa'].contacto_nombre}" if d['empresa'].contacto_nombre
                       else "Recibido por") + f"\nRepresentante legal · {d['empresa']}")


def documento_hallazgos(d):
    avisos = _avisos_comunes(d)
    if d["es_borrador"]:
        avisos.append(f"VERSIÓN PRELIMINAR: faltan revisar {d['pendientes_modelo']} atípicas y {d['pendientes_reglas']} "
                      f"avisos; {len(d.get('en_seguimiento', []))} atípicas en seguimiento (no incluidas).")
    o = d["por_origen"]
    datos = _datos_comunes(d) + [
        ("Hallazgos", f"{len(d['hallazgos'])} (modelo {o['modelo']}, reglas {o['reglas']}, ambos {o['ambos']})"),
        ("Importe de los hallazgos", f"{importe(d['importe_total'])} ({porcentaje(d['porcentaje_importe'], 2)} "
                                     f"del movimiento total)"),
        ("En cuentas tributarias", f"{d['tributarios']['cantidad']} por {importe(d['tributarios']['importe'])}"),
    ]
    secciones = []
    alcance = _seccion_alcance(d)
    if alcance:
        secciones.append(alcance)
    columnas = ["N.º", "Fecha", "Fila", "Comprobante", "Cuenta", "Glosa", "Importe", "Origen",
                "Qué se detectó", "Fundamento del auditor"]
    anchos = [0.5, 1.0, 0.65, 1.3, 2.1, 2.3, 1.35, 1.1, 3.1, 2.9]
    if d["varias_cargas"]:
        columnas.insert(1, "Carga")
        anchos.insert(1, 0.6)
    filas = []
    for h in d["hallazgos"]:
        r = h["registro"]
        detectado = []
        fundamento = []
        if h["modelo"]:
            detectado += [("También: " if i else "") + m["texto"] for i, m in enumerate(h["modelo"]["motivos"])]
            if h["modelo"]["comentario"]:
                fundamento.append(h["modelo"]["comentario"])
        for regla in h["reglas"]:
            detectado.append(("Regla: " if h["modelo"] else "") + regla["descripcion"])
            if regla["comentario"]:
                fundamento.append(regla["comentario"])
        fila = [h["numero"], fecha(r.fecha) if r else "", r.fila_origen if r else "",
                r.numero_comprobante if r else "",
                (f"{r.cuenta.codigo} {r.cuenta.nombre}" + (" [tributaria]" if h["tributaria"] else "")) if r else "",
                r.glosa if r else "", _importe_lado(r), h["origen"], "\n".join(detectado), "\n".join(fundamento)]
        if d["varias_cargas"]:
            fila.insert(1, h["carga_id"])
        filas.append(fila)
    secciones.append(Seccion("Detalle de los hallazgos (de mayor a menor importe)", columnas, filas,
                             anchos=anchos, numericas=(columnas.index("Importe"),), hoja="Hallazgos"))
    if d.get("en_seguimiento"):
        secciones.append(Seccion(
            "En seguimiento: atípicas que requieren revisar otros documentos o cuentas",
            ["Carga", "Fecha", "Fila", "Cuenta", "Glosa", "Importe", "Qué falta revisar", "Anotado por"],
            [[s.analisis.carga_id, fecha(s.registro.fecha), s.registro.fila_origen,
              f"{s.registro.cuenta.codigo} {s.registro.cuenta.nombre}", s.registro.glosa,
              _importe_lado(s.registro), s.comentario_revision,
              f"{s.revisado_por.username if s.revisado_por else ''} {fecha(s.fecha_revision)}".strip()]
             for s in d["en_seguimiento"]],
            anchos=[0.6, 1, 0.6, 2.2, 2.6, 1.3, 3.2, 1.5], numericas=(5,), hoja="En seguimiento",
            nota="No forman parte de los hallazgos hasta que el auditor las marque como observadas o válidas.",
            # Puntos abiertos: el papel de trabajo impreso o en PDF lleva solo hallazgos.
            solo_excel=True))
    return Documento(
        tipo="hallazgos", titulo="Hallazgos de auditoría",
        subtitulo=("Transacciones confirmadas como hallazgo por el auditor: atípicas del modelo marcadas "
                   "como observadas y avisos de las reglas marcados como observados."),
        datos=datos, secciones=secciones, avisos=avisos, emision=d["emision"].id,
        nombre_archivo=_nombre_archivo("hallazgos", d["empresa"], d["gestion"], d.get("carga")))


def documento_expediente(empresa, filas, usuario):
    """Comparativo de las gestiones de una empresa."""
    def var(valor):
        return "" if valor is None else ("+" if valor > 0 else "") + porcentaje(valor, 0)
    tabla = [[f["gestion"].anio, f"{fecha(f['gestion'].fecha_inicio)} al {fecha(f['gestion'].fecha_fin)}",
              f["cargas"], f["anuladas"], f["registros"], f["atipicas"], var(f["variacion_atipicas"]),
              importe(f["importe_atipicas"]), var(f["variacion_importe_atipicas"]),
              porcentaje(f["porcentaje_importe"]), f["revisadas"], f["relevantes"], f["observaciones"],
              var(f["variacion_observaciones"])] for f in filas]
    return Documento(
        tipo="expediente", titulo="Comparativo de gestiones",
        subtitulo="Cifras de cada gestión de la empresa, con su variación frente a la gestión anterior.",
        datos=[("Empresa auditada", f"{empresa}" + (f" (NIT {empresa.nit})" if empresa.nit else "")),
               ("Emitido", f"{fecha(timezone.now(), True)} por {usuario.username}")],
        secciones=[Seccion("Gestiones",
            ["Gestión", "Período", "Cargas válidas", "Anuladas", "Registros", "Atípicas", "Var. atípicas",
             "Importe atípico", "Var. importe", "% del movimiento", "Revisadas", "Útiles",
             "Observaciones", "Var. observ."],
            tabla, anchos=[0.8, 2, 0.8, 0.8, 1, 0.9, 0.9, 1.5, 0.9, 1, 0.9, 0.9, 1, 0.9],
            numericas=(4, 5, 7, 10, 11, 12), hoja="Gestiones")],
        nombre_archivo=f"comparativo_{slugify(empresa.nombre)[:40]}_{timezone.localtime().strftime('%Y%m%d_%H%M')}")


# Excel

_INICIO_FORMULA = re.compile(r"^[=+\-@\t\r]")


def _celda(hoja, fila, columna, valor):
    """Escribe un valor; los textos que empiezan como fórmula se guardan como texto
    (protección contra inyección de fórmulas en Excel)."""
    celda = hoja.cell(row=fila, column=columna, value=valor)
    if isinstance(valor, str) and _INICIO_FORMULA.match(valor):
        celda.data_type = "s"
    return celda


def a_excel(doc):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    guindo = "6E1B3F"
    cabecera = PatternFill("solid", fgColor=guindo)
    suave = PatternFill("solid", fgColor="F8F2F5")
    borde = Border(*(Side(style="thin", color="D9C9D1"),) * 4)
    libro = Workbook()

    portada = libro.active
    portada.title = "Resumen"
    portada["A1"] = RAZON_SOCIAL
    portada["A1"].font = Font(bold=True, size=12, color=guindo)
    portada["A2"] = doc.titulo
    portada["A2"].font = Font(bold=True, size=14)
    _celda(portada, 3, 1, doc.subtitulo).alignment = Alignment(wrap_text=True, vertical="top")
    portada.merge_cells("A3:D3")
    portada.row_dimensions[3].height = 45
    fila = 5
    for aviso in doc.avisos:
        _celda(portada, fila, 1, aviso).font = Font(bold=True, color="9C5700")
        fila += 1
    if doc.avisos:
        fila += 1
    for etiqueta, valor in doc.datos:
        _celda(portada, fila, 1, etiqueta).font = Font(bold=True)
        _celda(portada, fila, 2, valor)
        fila += 1
    _celda(portada, fila + 1, 1, PIE).font = Font(italic=True, size=8, color="777777")
    portada.column_dimensions["A"].width = 30
    portada.column_dimensions["B"].width = 80

    usados = set()
    for seccion in doc.secciones:
        nombre = re.sub(r"[\[\]:*?/\\]", "", seccion.hoja or seccion.titulo)[:28].strip() or "Hoja"
        base, n = nombre, 2
        while nombre in usados:
            nombre = f"{base[:25]} {n}"
            n += 1
        usados.add(nombre)
        hoja = libro.create_sheet(nombre)
        hoja["A1"] = seccion.titulo
        hoja["A1"].font = Font(bold=True, size=12, color=guindo)
        for c, titulo in enumerate(seccion.columnas, start=1):
            celda = hoja.cell(row=3, column=c, value=titulo)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = cabecera
            celda.alignment = Alignment(wrap_text=True, vertical="center")
            celda.border = borde
        for f, valores in enumerate(seccion.filas, start=4):
            for c, valor in enumerate(valores, start=1):
                celda = _celda(hoja, f, c, valor)
                celda.border = borde
                celda.alignment = Alignment(wrap_text=True, vertical="top",
                                            horizontal="right" if (c - 1) in seccion.numericas else "left")
                if f % 2 == 1:
                    celda.fill = suave
        if seccion.nota:
            _celda(hoja, len(seccion.filas) + 5, 1, seccion.nota).font = Font(italic=True, size=9)
        anchos = seccion.anchos or [1] * len(seccion.columnas)
        for c, proporcion in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(c)].width = max(9, min(60, proporcion * 11))
        hoja.freeze_panes = "A4"
        if seccion.filas:
            hoja.auto_filter.ref = f"A3:{get_column_letter(len(seccion.columnas))}{len(seccion.filas) + 3}"
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.fitToWidth = 1
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.page_setup.fitToHeight = 0

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# PDF

def a_pdf(doc):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table,
                                    TableStyle)
    from xml.sax.saxutils import escape

    guindo = colors.HexColor("#6E1B3F")
    oscuro = colors.HexColor("#4A1129")
    borde = colors.HexColor("#D9C9D1")
    suave = colors.HexColor("#F8F2F5")
    estilos = getSampleStyleSheet()
    normal = ParagraphStyle("normal", parent=estilos["Normal"], fontName="Helvetica", fontSize=7.5, leading=9)
    derecha = ParagraphStyle("derecha", parent=normal, alignment=TA_RIGHT)
    cabeza = ParagraphStyle("cabeza", parent=normal, fontName="Helvetica-Bold", textColor=colors.white)
    titulo = ParagraphStyle("titulo", parent=estilos["Title"], fontName="Times-Bold", fontSize=16,
                            textColor=oscuro, alignment=0, spaceAfter=4)
    subtitulo = ParagraphStyle("subtitulo", parent=normal, fontSize=8.5, leading=11, textColor=colors.HexColor("#555555"))
    seccion_estilo = ParagraphStyle("seccion", parent=estilos["Heading3"], fontName="Times-Bold", fontSize=11,
                                    textColor=oscuro, spaceBefore=10, spaceAfter=4)
    aviso_estilo = ParagraphStyle("aviso", parent=normal, fontSize=8.5, textColor=colors.HexColor("#8A4B00"),
                                  backColor=colors.HexColor("#FFF4DE"), borderPadding=4, spaceAfter=6)

    def p(texto, estilo=normal):
        return Paragraph(escape(str(texto if texto is not None else "")).replace("\n", "<br/>"), estilo)

    salida = io.BytesIO()
    pagina = landscape(letter)
    ancho_util = pagina[0] - 24 * mm

    def pie(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#777777"))
        texto = f"{RAZON_SOCIAL} · {doc.titulo}"
        if doc.emision:
            texto += f" · Emisión N.º {doc.emision}"
        canvas.drawString(12 * mm, 8 * mm, texto)
        canvas.drawRightString(pagina[0] - 12 * mm, 8 * mm, f"Página {documento.page}")
        canvas.restoreState()

    plantilla = SimpleDocTemplate(salida, pagesize=pagina, leftMargin=12 * mm, rightMargin=12 * mm,
                                  topMargin=12 * mm, bottomMargin=14 * mm, title=doc.titulo,
                                  author=RAZON_SOCIAL)
    elementos = []

    logo = Path(settings.BASE_DIR) / "static" / "img" / "logo-sts-claro.png"
    membrete = [p(RAZON_SOCIAL, ParagraphStyle("rs", parent=normal, fontName="Times-Bold", fontSize=11,
                                              leading=14, textColor=oscuro)),
                p("Auditoría externa e informática forense", subtitulo)]
    if logo.exists():
        try:
            imagen = Image(str(logo), width=38 * mm, height=13 * mm, kind="proportional")
            etiqueta = {"hallazgos": "PAPEL DE TRABAJO",
                        "observaciones": "PARA FIRMAR CON LA EMPRESA"}.get(doc.tipo, "REPORTE DE APOYO")
            encabezado = Table([[imagen, membrete, p(etiqueta, ParagraphStyle(
                "pt", parent=normal, fontName="Helvetica-Bold", alignment=TA_RIGHT, textColor=oscuro))]],
                colWidths=[42 * mm, ancho_util - 92 * mm, 50 * mm])
        except Exception:
            encabezado = Table([[membrete]], colWidths=[ancho_util])
    else:
        encabezado = Table([[membrete]], colWidths=[ancho_util])
    encabezado.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                    ("LINEBELOW", (0, 0), (-1, 0), 2, guindo),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    elementos += [encabezado, Spacer(1, 6)]
    for aviso in doc.avisos:
        elementos.append(p(aviso, aviso_estilo))
    elementos += [p(doc.titulo, titulo), p(doc.subtitulo, subtitulo), Spacer(1, 6)]

    datos = Table([[p(e, ParagraphStyle("e", parent=normal, fontName="Helvetica-Bold")), p(v)] for e, v in doc.datos],
                  colWidths=[45 * mm, ancho_util - 45 * mm])
    datos.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), suave), ("BOX", (0, 0), (-1, -1), 0.5, borde),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2),
                               ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
    elementos.append(datos)

    for seccion in doc.secciones:
        if seccion.solo_excel:
            continue
        bloque = [p(seccion.titulo, seccion_estilo)]
        if seccion.filas:
            proporciones = seccion.anchos or [1] * len(seccion.columnas)
            total = sum(proporciones)
            anchos = [ancho_util * x / total for x in proporciones]
            filas = [[p(c, cabeza) for c in seccion.columnas]]
            for valores in seccion.filas:
                filas.append([p(v, derecha if i in seccion.numericas else normal) for i, v in enumerate(valores)])
            tabla = Table(filas, colWidths=anchos, repeatRows=1)
            estilo = [("BACKGROUND", (0, 0), (-1, 0), oscuro), ("GRID", (0, 0), (-1, -1), 0.4, borde),
                      ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 2),
                      ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
            for i in range(2, len(filas), 2):
                estilo.append(("BACKGROUND", (0, i), (-1, i), suave))
            tabla.setStyle(TableStyle(estilo))
            bloque.append(tabla)
        else:
            bloque.append(p("Sin datos.", subtitulo))
        if seccion.nota:
            bloque += [Spacer(1, 2), p(seccion.nota, subtitulo)]
        # Las tablas cortas no se parten entre páginas; las largas sí (repiten la cabecera).
        elementos.append(KeepTogether(bloque) if len(seccion.filas) <= 15 else bloque[0])
        if len(seccion.filas) > 15:
            elementos += bloque[1:]

    # Solo Hallazgos y Observaciones llevan firmas. En Hallazgos, el Administrador
    # firma una vez si es el emisor; si emite un auditor, el socio debe aprobar.
    if doc.tipo in ("hallazgos", "observaciones"):
        firma = ParagraphStyle("firma", parent=normal, fontSize=8, leading=10, alignment=1)  # centrado
        if doc.tipo == "observaciones":
            cargo = "Socio de auditoría" if doc.emisor_es_socio else "Auditor"
            bloques = [f"Elaborado por: {escape(doc.emitido_por)}<br/>{cargo} · {escape(RAZON_SOCIAL)}",
                       (f"{escape(doc.contador)}<br/>Contador de la empresa" if doc.contador
                        else "Contador de la empresa<br/>Nombre: ______________________"),
                       escape(doc.firma_cliente).replace("\n", "<br/>")]
        elif doc.emisor_es_socio:
            bloques = [f"Elaborado y aprobado por: {escape(doc.emitido_por)}<br/>Socio de auditoría · Fecha: ____/____/______"]
        else:
            bloques = [f"Elaborado por: {escape(doc.emitido_por)}<br/>Auditor · Fecha: ____/____/______",
                       "Revisado y aprobado por<br/>Socio de auditoría · Fecha: ____/____/______"]
        celdas = [Paragraph(b, firma) for b in bloques]
        firmas = Table([celdas], colWidths=[ancho_util / max(2, len(celdas))] * len(celdas), hAlign="LEFT")
        firmas.setStyle(TableStyle([("LINEABOVE", (i, 0), (i, 0), 0.6, colors.black) for i in range(len(celdas))]
                                   + [("LEFTPADDING", (0, 0), (-1, -1), 18), ("RIGHTPADDING", (0, 0), (-1, -1), 18)]))
        elementos += [Spacer(1, 34), KeepTogether([firmas, Spacer(1, 10), p(PIE, subtitulo)])]
    else:
        elementos += [Spacer(1, 10), p(PIE, subtitulo)]

    plantilla.build(elementos, onFirstPage=pie, onLaterPages=pie)
    return salida.getvalue()
